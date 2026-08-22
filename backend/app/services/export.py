"""导出能力：Markdown / PDF / Anki 卡片。"""
import csv
import json
import os
import tempfile
from io import BytesIO, StringIO
from typing import List, Optional

from ..utils.timestamp import normalize_hms


def render_note_markdown(note: dict, subject: str = "general", words: Optional[dict] = None, meta: Optional[dict] = None, frames: Optional[dict] = None) -> str:
    """生成排版美观的笔记型 Markdown 文档。frames: {章节索引: [相对图片路径]}"""
    fence = chr(96) * 3
    tick = chr(96)
    title = str(note.get("title") or "课程笔记")
    lines = ["# " + title, ""]
    summary = str(note.get("summary") or "").strip()
    if summary:
        lines.append("> " + summary)
        lines.append("")
    meta = meta or {}
    meta_parts = []
    if meta.get("bvid"):
        meta_parts.append("**视频**：B站 " + str(meta["bvid"]) + " P" + str(meta.get("page") or 1))
    if meta.get("subject_name"):
        meta_parts.append("**学科**：" + str(meta["subject_name"]))
    if meta.get("created_at"):
        meta_parts.append("**生成时间**：" + str(meta["created_at"])[:19].replace("T", " "))
    if meta_parts:
        lines.append(" · ".join(meta_parts))
        lines.append("")
    chapters = note.get("chapters") or []
    if chapters:
        lines.append("## 目录")
        lines.append("")
        for idx, ch in enumerate(chapters, start=1):
            lines.append(str(idx) + ". " + str(ch.get("title") or "章节"))
        lines.append("")
    for idx, ch in enumerate(chapters, start=1):
        lines.append("---")
        lines.append("")
        lines.append("## " + str(idx) + ". " + str(ch.get("title") or "章节"))
        lines.append("")
        if frames and (idx - 1) in frames:
            for img in frames[idx - 1]:
                lines.append("![" + str(img) + "](" + str(img) + ")")
            lines.append("")
        for p in ch.get("points") or []:
            if not isinstance(p, dict):
                continue
            content = str(p.get("content") or "").strip()
            if not content:
                continue
            ts = str(p.get("time_stamp") or "")
            prefix = ("**" + tick + ts + tick + "** ") if ts else ""
            mark = " **【重点】**" if p.get("important") else ""
            lines.append("- " + prefix + content + mark)
        lines.append("")
    mindmap = str(note.get("mindmap") or "").strip()
    if mindmap:
        lines.append("---")
        lines.append("")
        lines.append("## 知识框架（思维导图）")
        lines.append("")
        lines.append(fence + "mermaid")
        lines.append(mindmap)
        lines.append(fence)
        lines.append("")
    if words:
        word_list = words.get("words") or []
        if word_list:
            lines.append("---")
            lines.append("")
            lines.append("## 生词本")
            lines.append("")
            lines.append("| 单词 | 音标 | 释义 | 原句 | 时间戳 |")
            lines.append("| --- | --- | --- | --- | --- |")
            for w in word_list:
                lines.append(
                    "| " + _cell(w.get("word")) + " | " + _cell(w.get("phonetic")) + " | "
                    + _cell(w.get("meaning")) + " | " + _cell(w.get("sentence")) + " | "
                    + _cell(w.get("time_stamp")) + " |"
                )
            lines.append("")
        pron = words.get("pronunciations") or []
        if pron:
            lines.append("## 语音现象标注")
            lines.append("")
            for it in pron:
                lines.append(
                    "- 【" + str(it.get("phenomenon") or "") + "】 "
                    + str(it.get("sentence") or "") + "（" + str(it.get("time_stamp") or "") + "）"
                )
            lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("> 本笔记由 BiliLearn-AI 自动生成，时间戳可在网页端点击跳转到视频对应片段。")
    return "\n".join(lines)


def _cell(value) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", " ")


def markdown_to_pdf_bytes(markdown: str, title: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("未安装 reportlab，无法导出 PDF") from exc

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, title=title or "BiliLearn 笔记",
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = {
        "title": ParagraphStyle("title", fontName="STSong-Light", fontSize=18, leading=26, spaceAfter=12),
        "h2": ParagraphStyle("h2", fontName="STSong-Light", fontSize=14, leading=20, spaceBefore=10, spaceAfter=6),
        "body": ParagraphStyle("body", fontName="STSong-Light", fontSize=10.5, leading=17),
    }
    story = []
    for line in markdown.split("\n"):
        line = line.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            story.append(Paragraph(_escape(line[2:]), styles["title"]))
        elif line.startswith("## "):
            story.append(Paragraph(_escape(line[3:]), styles["h2"]))
        elif line.startswith("### "):
            story.append(Paragraph(_escape(line[4:]), styles["h2"]))
        elif line.strip().startswith(chr(96) * 3):
            continue
        else:
            story.append(Paragraph(_escape(line), styles["body"]))
    doc.build(story)
    return buf.getvalue()


def build_note_pdf(note: dict, subject: str = "general", words: Optional[dict] = None, meta: Optional[dict] = None, frames: Optional[dict] = None) -> bytes:
    """基于结构化笔记生成排版良好的 PDF 文档。frames: {章节索引: [图片路径]}"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import HRFlowable, Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("未安装 reportlab，无法导出 PDF") from exc

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buf = BytesIO()
    title = str(note.get("title") or "课程笔记")
    meta = meta or {}
    C_MAIN = colors.HexColor("#1f2d3d")
    C_SUB = colors.HexColor("#909399")
    C_BLUE = colors.HexColor("#409eff")
    C_BG = colors.HexColor("#e8f4ff")
    C_LINE = colors.HexColor("#e4e7ed")

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(C_SUB)
        canvas.drawString(1.6 * cm, 1.0 * cm, "BiliLearn-AI 学习笔记")
        canvas.drawRightString(A4[0] - 1.6 * cm, 1.0 * cm, "第 " + str(canvas.getPageNumber()) + " 页")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=A4, title=title,
        leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.8 * cm, bottomMargin=2.0 * cm,
    )
    styles = {
        "title": ParagraphStyle("t", fontName="STSong-Light", fontSize=20, leading=28, spaceAfter=8, textColor=C_MAIN),
        "meta": ParagraphStyle("m", fontName="STSong-Light", fontSize=9, leading=14, textColor=C_SUB, spaceAfter=10),
        "quote": ParagraphStyle("q", fontName="STSong-Light", fontSize=10, leading=16, textColor=colors.HexColor("#606266"), leftIndent=8, borderPadding=(6, 8, 6, 8), backColor=colors.HexColor("#f5f7fa")),
        "h1": ParagraphStyle("h1", fontName="STSong-Light", fontSize=14, leading=20, spaceBefore=16, spaceAfter=8, textColor=C_MAIN, backColor=C_BG, borderPadding=(4, 8, 4, 8)),
        "h2": ParagraphStyle("h2", fontName="STSong-Light", fontSize=12, leading=18, spaceBefore=10, spaceAfter=6, textColor=C_MAIN),
        "bullet": ParagraphStyle("b", fontName="STSong-Light", fontSize=10.5, leading=17, leftIndent=16, bulletIndent=4, spaceAfter=5),
    }
    story = [Paragraph(_escape(title), styles["title"])]
    summary = str(note.get("summary") or "").strip()
    if summary:
        story.append(Paragraph(_escape(summary), styles["quote"]))
    meta_parts = []
    if meta.get("bvid"):
        meta_parts.append("视频：B站 " + str(meta["bvid"]) + " P" + str(meta.get("page") or 1))
    if meta.get("subject_name"):
        meta_parts.append("学科：" + str(meta["subject_name"]))
    if meta.get("created_at"):
        meta_parts.append("生成时间：" + str(meta["created_at"])[:19].replace("T", " "))
    if meta_parts:
        story.append(Paragraph(_escape(" ｜ ".join(meta_parts)), styles["meta"]))
    story.append(HRFlowable(width="100%", thickness=1, color=C_LINE, spaceAfter=6))

    chapters = note.get("chapters") or []
    for idx, ch in enumerate(chapters, start=1):
        story.append(Paragraph(_escape(str(idx) + "、" + str(ch.get("title") or "章节")), styles["h1"]))
        if frames and (idx - 1) in frames:
            for img in frames[idx - 1]:
                try:
                    story.append(RLImage(img, width=14 * cm))
                    story.append(Spacer(1, 4))
                except Exception:
                    pass
        for p in ch.get("points") or []:
            if not isinstance(p, dict):
                continue
            content = str(p.get("content") or "").strip()
            if not content:
                continue
            ts = str(p.get("time_stamp") or "")
            prefix = ("<b>[" + ts + "]</b> ") if ts else ""
            mark = " <b>【重点】</b>" if p.get("important") else ""
            story.append(Paragraph(prefix + _escape(content) + mark, styles["bullet"], bulletText="•"))

    word_list = (words or {}).get("words") or []
    if word_list:
        story.append(Paragraph(_escape("生词本"), styles["h2"]))
        data = [["单词", "音标", "释义", "原句", "时间戳"]]
        for w in word_list[:50]:
            data.append([
                _escape(str(w.get("word") or "")),
                _escape(str(w.get("phonetic") or "")),
                _escape(str(w.get("meaning") or "")),
                _escape(str(w.get("sentence") or "")),
                _escape(str(w.get("time_stamp") or "")),
            ])
        tbl = Table(data, colWidths=[2.4 * cm, 2.8 * cm, 4.2 * cm, 5.4 * cm, 2.0 * cm])
        tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), C_BG),
            ("TEXTCOLOR", (0, 0), (-1, 0), C_MAIN),
            ("GRID", (0, 0), (-1, -1), 0.5, C_LINE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(Spacer(1, 4))
        story.append(tbl)

    pron = (words or {}).get("pronunciations") or []
    if pron:
        story.append(Paragraph(_escape("语音现象标注"), styles["h2"]))
        for it in pron:
            story.append(Paragraph(
                _escape("【" + str(it.get("phenomenon") or "") + "】" + str(it.get("sentence") or "")
                        + "（" + str(it.get("time_stamp") or "") + "）"),
                styles["bullet"], bulletText="•",
            ))
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_note_docx(note: dict, subject: str = "general", words: Optional[dict] = None, meta: Optional[dict] = None, frames: Optional[dict] = None) -> bytes:
    """生成排版美观的 Word 笔记（含可选关键帧截图）。frames: {章节索引: [图片路径]}"""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("未安装 python-docx，无法导出 Word") from exc

    C_MAIN = RGBColor(0x1F, 0x2D, 0x3D)
    C_BLUE = RGBColor(0x33, 0x70, 0xFF)
    C_RED = RGBColor(0xF5, 0x6C, 0x6C)
    C_GRAY = RGBColor(0x90, 0x99, 0x99)

    def _set_font(obj, size=None, bold=None, color=None, name="微软雅黑"):
        if size is not None:
            obj.font.size = Pt(size)
        if bold is not None:
            obj.font.bold = bold
        if color is not None:
            obj.font.color.rgb = color
        obj.font.name = name
        try:
            rpr = obj._element.get_or_add_rPr()
            rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)
        except Exception:
            pass

    doc = Document()
    try:
        _set_font(doc.styles["Normal"], size=10.5)
        for sname in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            try:
                _set_font(doc.styles[sname], name="微软雅黑")
            except Exception:
                pass
    except Exception:
        pass

    title = str(note.get("title") or "课程笔记")
    p = doc.add_paragraph()
    _set_font(p.add_run(title), size=20, bold=True, color=C_MAIN)
    summary = str(note.get("summary") or "").strip()
    if summary:
        p = doc.add_paragraph()
        _set_font(p.add_run(summary), size=10.5, color=RGBColor(0x60, 0x62, 0x66))
        p.paragraph_format.space_after = Pt(6)
    meta = meta or {}
    meta_parts = []
    if meta.get("bvid"):
        meta_parts.append("视频：B站 " + str(meta["bvid"]) + " P" + str(meta.get("page") or 1))
    if meta.get("subject_name"):
        meta_parts.append("学科：" + str(meta["subject_name"]))
    if meta.get("created_at"):
        meta_parts.append("生成时间：" + str(meta["created_at"])[:19].replace("T", " "))
    if meta_parts:
        p = doc.add_paragraph()
        _set_font(p.add_run(" ｜ ".join(meta_parts)), size=9, color=C_GRAY)

    chapters = note.get("chapters") or []
    for ci, ch in enumerate(chapters, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        _set_font(p.add_run(str(ci) + "、" + str(ch.get("title") or "章节")), size=14, bold=True, color=C_MAIN)
        if frames and ci - 1 in frames:
            for img in frames[ci - 1]:
                try:
                    doc.add_picture(img, width=Inches(5.6))
                except Exception:
                    pass
        for pt in ch.get("points") or []:
            if not isinstance(pt, dict):
                continue
            content = str(pt.get("content") or "").strip()
            if not content:
                continue
            p = doc.add_paragraph(style="List Bullet")
            ts = str(pt.get("time_stamp") or "")
            if ts:
                _set_font(p.add_run("[" + ts + "] "), size=9.5, bold=True, color=C_BLUE)
            _set_font(p.add_run(content), size=10.5)
            if pt.get("important"):
                _set_font(p.add_run(" 【重点】"), size=10.5, bold=True, color=C_RED)

    word_list = (words or {}).get("words") or []
    if word_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        _set_font(p.add_run("生词本"), size=14, bold=True, color=C_MAIN)
        table = doc.add_table(rows=1, cols=5)
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        for i, h in enumerate(["单词", "音标", "释义", "原句", "时间戳"]):
            _set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=10, bold=True)
        for w in word_list[:50]:
            cells = table.add_row().cells
            for i, key in enumerate(["word", "phonetic", "meaning", "sentence", "time_stamp"]):
                _set_font(cells[i].paragraphs[0].add_run(str(w.get(key) or "")), size=9.5)

    pron = (words or {}).get("pronunciations") or []
    if pron:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        _set_font(p.add_run("语音现象标注"), size=14, bold=True, color=C_MAIN)
        for it in pron:
            p = doc.add_paragraph(style="List Bullet")
            _set_font(p.add_run("【" + str(it.get("phenomenon") or "") + "】" + str(it.get("sentence") or "")), size=10.5)
            if it.get("time_stamp"):
                _set_font(p.add_run("（" + str(it["time_stamp"]) + "）"), size=9.5, color=C_GRAY)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_note_xmind(note: dict) -> bytes:
    """生成 .xmind 思维导图文件（XMind 2020+ 兼容 zip 格式）。"""
    import json
    import uuid
    import zipfile

    def _node(title, children=None):
        n = {"id": "n" + uuid.uuid4().hex[:12], "class": "topic", "title": str(title)}
        if children:
            n["children"] = {"attached": children}
        return n

    root_children = []
    for ch in note.get("chapters") or []:
        pts = (ch.get("points") or [])[:8]
        children = []
        for p in pts:
            label = str(p.get("content") or "")[:80]
            ts = str(p.get("time_stamp") or "")
            if ts:
                label += "  " + ts
            children.append(_node(label))
        root_children.append(_node(ch.get("title") or "章节", children))
    root = _node(note.get("title") or "课程笔记", root_children)
    sheet = {
        "id": "s" + uuid.uuid4().hex[:12],
        "class": "sheet",
        "title": "BiliLearn 笔记",
        "rootTopic": root,
    }
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("content.json", json.dumps([sheet], ensure_ascii=False))
        z.writestr("manifest.json", json.dumps({"file-entries": {"content.json": {}, "metadata.json": {}}}))
        z.writestr("metadata.json", json.dumps({"creator": {"name": "BiliLearn-AI", "version": "1.2"}}))
    return buf.getvalue()


def build_anki_apkg(words: List[dict], deck_name: str = "BiliLearn 生词本") -> bytes:
    try:
        import genanki
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("未安装 genanki，无法导出 Anki 卡片") from exc
    model = genanki.Model(
        1607392319,
        "BiliLearn Word",
        fields=[
            {"name": "word"}, {"name": "phonetic"}, {"name": "meaning"},
            {"name": "sentence"}, {"name": "time"},
        ],
        templates=[{
            "name": "单词卡",
            "qfmt": "{{word}}<br/><small>{{phonetic}}</small>",
            "afmt": '{{FrontSide}}<hr id="answer">{{meaning}}<br/><i>{{sentence}}</i><br/><small>{{time}}</small>',
        }],
    )
    deck = genanki.Deck(2059400110, deck_name)
    for w in words or []:
        deck.add_note(genanki.Note(model=model, fields=[
            str(w.get("word") or ""),
            str(w.get("phonetic") or ""),
            str(w.get("meaning") or ""),
            str(w.get("sentence") or ""),
            str(normalize_hms(w.get("time_stamp")) if w.get("time_stamp") else ""),
        ]))
    pkg = genanki.Package(deck)
    fd, path = tempfile.mkstemp(suffix=".apkg")
    os.close(fd)
    try:
        pkg.write_to_file(path)
        with open(path, "rb") as f:
            return f.read()
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def words_to_csv(words: List[dict]) -> str:
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(["word", "phonetic", "meaning", "sentence", "time_stamp"])
    for w in words or []:
        writer.writerow([
            w.get("word", ""), w.get("phonetic", ""), w.get("meaning", ""),
            w.get("sentence", ""), w.get("time_stamp", ""),
        ])
    return "\ufeff" + buf.getvalue()
